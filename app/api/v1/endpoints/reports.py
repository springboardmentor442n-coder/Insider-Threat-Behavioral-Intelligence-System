import io
import csv
from datetime import datetime, timezone
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse, HTMLResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc

from app.core.database import get_db
from app.core.security import require_analyst
from app.models import Employee, Alert, Incident, RiskScore, Department

router = APIRouter(prefix="/reports", tags=["Reports & Exports"])


@router.get("/export/csv")
def export_csv(
    report_type: str = Query("incidents", regex="^(incidents|alerts|employees)$"),
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    """
    Export incident, alert, or employee data as a CSV spreadsheet.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    filename = f"ueba_report_{report_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"

    if report_type == "incidents":
        # Header
        writer.writerow(["Incident ID", "Title", "Status", "Severity", "Impact Level", "Employee", "Assigned Analyst", "Created At"])
        incidents = db.query(Incident).order_by(desc(Incident.created_at)).all()
        for inc in incidents:
            emp_name = inc.employee.full_name if inc.employee else "N/A"
            analyst_name = inc.assigned_analyst.full_name if inc.assigned_analyst else "Unassigned"
            writer.writerow([
                inc.incident_id,
                inc.title,
                inc.status.value if inc.status else "",
                inc.severity.value if inc.severity else "",
                inc.impact_level,
                emp_name,
                analyst_name,
                inc.created_at.strftime("%Y-%m-%d %H:%M:%S") if inc.created_at else ""
            ])

    elif report_type == "alerts":
        writer.writerow(["Alert ID", "Title", "Severity", "Status", "Employee", "Triggered At", "Notes"])
        alerts = db.query(Alert).order_by(desc(Alert.triggered_at)).all()
        for alert in alerts:
            emp_name = alert.employee.full_name if alert.employee else "N/A"
            writer.writerow([
                alert.alert_id,
                alert.title,
                alert.severity.value if alert.severity else "",
                alert.status.value if alert.status else "",
                emp_name,
                alert.triggered_at.strftime("%Y-%m-%d %H:%M:%S") if alert.triggered_at else "",
                alert.notes or ""
            ])

    elif report_type == "employees":
        writer.writerow(["Employee ID", "Full Name", "Email", "Designation", "Department", "Current Risk Score", "Risk Category"])
        # Query latest risk scores
        from app.api.v1.endpoints.security import get_risk_leaderboard
        leaderboard = get_risk_leaderboard(db, limit=100)
        for entry in leaderboard:
            emp = db.query(Employee).filter(Employee.id == entry.employee_id).first()
            dept_name = emp.department.name if emp and emp.department else "N/A"
            writer.writerow([
                emp.employee_id if emp else "N/A",
                entry.employee_name,
                emp.email if emp else "N/A",
                emp.designation if emp else "N/A",
                dept_name,
                entry.current_score,
                entry.risk_category.value if entry.risk_category else ""
            ])

    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/export/pdf", response_class=HTMLResponse)
def export_pdf(
    report_type: str = Query("incidents", regex="^(incidents|alerts|employees)$"),
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    """
    Returns a print-friendly HTML document styled for high-quality browser Save-as-PDF.
    """
    title = f"Insider Threat Report - {report_type.upper()}"
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    # Construct tables based on type
    rows_html = ""
    headers = []

    if report_type == "incidents":
        headers = ["ID", "Title", "Status", "Severity", "Impact", "Employee", "Created At"]
        items = db.query(Incident).order_by(desc(Incident.created_at)).all()
        for item in items:
            emp = item.employee.full_name if item.employee else "N/A"
            rows_html += f"""
            <tr>
                <td class="mono">{item.incident_id}</td>
                <td><strong>{item.title}</strong></td>
                <td><span class="badge badge-{item.status.value}">{item.status.value}</span></td>
                <td><span class="badge badge-{item.severity.value}">{item.severity.value}</span></td>
                <td>{item.impact_level}</td>
                <td>{emp}</td>
                <td>{item.created_at.strftime('%Y-%m-%d') if item.created_at else ''}</td>
            </tr>
            """
    elif report_type == "alerts":
        headers = ["ID", "Alert Title", "Severity", "Status", "Employee", "Triggered At"]
        items = db.query(Alert).order_by(desc(Alert.triggered_at)).all()
        for item in items:
            emp = item.employee.full_name if item.employee else "N/A"
            rows_html += f"""
            <tr>
                <td class="mono">{item.alert_id}</td>
                <td>{item.title}</td>
                <td><span class="badge badge-{item.severity.value}">{item.severity.value}</span></td>
                <td><span class="badge badge-{item.status.value}">{item.status.value}</span></td>
                <td>{emp}</td>
                <td>{item.triggered_at.strftime('%Y-%m-%d %H:%M') if item.triggered_at else ''}</td>
            </tr>
            """
    elif report_type == "employees":
        headers = ["ID", "Name", "Email", "Designation", "Department", "Risk Score", "Risk Category"]
        from app.api.v1.endpoints.security import get_risk_leaderboard
        leaderboard = get_risk_leaderboard(db, limit=100)
        for entry in leaderboard:
            emp = db.query(Employee).filter(Employee.id == entry.employee_id).first()
            dept = emp.department.name if emp and emp.department else "N/A"
            rows_html += f"""
            <tr>
                <td class="mono">{emp.employee_id if emp else 'N/A'}</td>
                <td><strong>{entry.employee_name}</strong></td>
                <td>{emp.email if emp else 'N/A'}</td>
                <td>{emp.designation if emp else 'N/A'}</td>
                <td>{dept}</td>
                <td class="mono">{entry.current_score}</td>
                <td><span class="badge badge-{entry.risk_category.value if entry.risk_category else 'low'}">{entry.risk_category.value if entry.risk_category else 'low'}</span></td>
            </tr>
            """

    headers_html = "".join(f"<th>{h}</th>" for h in headers)

    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <title>{title}</title>
        <style>
            @media print {{
                body {{ background: white; color: black; font-size: 11pt; }}
                .no-print {{ display: none; }}
                @page {{ margin: 1.5cm; }}
            }}
            body {{
                font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
                color: #2D3748;
                line-height: 1.5;
                margin: 0;
                padding: 40px;
                background-color: #F7FAFC;
            }}
            .container {{
                max-width: 1000px;
                margin: 0 auto;
                background: white;
                padding: 40px;
                border-radius: 8px;
                box-shadow: 0 4px 6px -1px rgba(0,0,0,0.1);
            }}
            .header {{
                border-bottom: 2px solid #E2E8F0;
                padding-bottom: 20px;
                margin-bottom: 30px;
                display: flex;
                justify-content: space-between;
                align-items: center;
            }}
            .logo {{
                font-weight: 800;
                font-size: 1.6rem;
                color: #1A202C;
                letter-spacing: -0.02em;
            }}
            .logo span {{ color: #3182CE; }}
            .title {{
                font-size: 1.8rem;
                font-weight: 800;
                margin: 0;
                color: #2D3748;
            }}
            .meta {{
                font-size: 0.85rem;
                color: #718096;
            }}
            table {{
                width: 100%;
                border-collapse: collapse;
                margin-top: 20px;
            }}
            th {{
                background-color: #EDF2F7;
                text-align: left;
                padding: 12px;
                font-weight: 700;
                font-size: 0.85rem;
                text-transform: uppercase;
                color: #4A5568;
                border-bottom: 2px solid #CBD5E0;
            }}
            td {{
                padding: 12px;
                border-bottom: 1px solid #E2E8F0;
                font-size: 0.9rem;
            }}
            .mono {{ font-family: monospace; font-weight: 600; }}
            .badge {{
                display: inline-block;
                padding: 2px 8px;
                border-radius: 4px;
                font-size: 0.75rem;
                font-weight: 700;
                text-transform: uppercase;
            }}
            .badge-low, .badge-resolved {{ background-color: #C6F6D5; color: #22543D; }}
            .badge-medium, .badge-investigating {{ background-color: #FEFCBF; color: #744210; }}
            .badge-high, .badge-open {{ background-color: #FEEBC8; color: #7B341E; }}
            .badge-critical {{ background-color: #FED7D7; color: #742A2A; }}
            .btn-print {{
                background-color: #3182CE;
                color: white;
                border: none;
                padding: 10px 20px;
                font-size: 0.9rem;
                font-weight: 600;
                border-radius: 6px;
                cursor: pointer;
                box-shadow: 0 2px 4px rgba(0,0,0,0.1);
            }}
            .btn-print:hover {{ background-color: #2B6CB0; }}
        </style>
        <script>
            window.onload = function() {{
                // Auto trigger print prompt for PDF export
                setTimeout(() => {{
                    window.print();
                }}, 500);
            }}
        </script>
    </head>
    <body>
        <div class="container">
            <div class="header">
                <div>
                    <div class="logo">ANTIGRAVITY <span>UEBA</span></div>
                    <h1 class="title">{title}</h1>
                    <div class="meta">Generated: {timestamp}</div>
                </div>
                <div class="no-print">
                    <button class="btn-print" onclick="window.print()">Print / Save PDF</button>
                </div>
            </div>
            <table>
                <thead>
                    <tr>{headers_html}</tr>
                </thead>
                <tbody>
                    {rows_html}
                </tbody>
            </table>
        </div>
    </body>
    </html>
    """
    return HTMLResponse(content=html_content)


@router.get("/export/excel")
def export_excel(
    report_type: str = Query("incidents", regex="^(incidents|alerts|employees)$"),
    db: Session = Depends(get_db),
    _=Depends(require_analyst),
):
    """
    Export incident, alert, or employee data as a formatted Excel spreadsheet.
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{report_type.capitalize()} Report"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    title_font = Font(name="Segoe UI", size=14, bold=True, color="1A365D")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    data_font = Font(name="Segoe UI", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    # Title Block
    ws.merge_cells("A1:G1")
    ws["A1"] = f"Insider Threat Detection System - {report_type.upper()} REPORT"
    ws["A1"].font = title_font
    ws["A1"].alignment = left_align
    ws.row_dimensions[1].height = 30
    
    ws["A2"] = f"Generated at: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    ws["A2"].font = Font(name="Segoe UI", size=9, italic=True)
    
    row_idx = 4
    
    if report_type == "incidents":
        headers = ["Incident ID", "Title", "Status", "Severity", "Impact Level", "Employee", "Assigned Analyst", "Created At"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
        ws.row_dimensions[row_idx].height = 24
        
        incidents = db.query(Incident).order_by(desc(Incident.created_at)).all()
        for inc in incidents:
            row_idx += 1
            emp_name = inc.employee.full_name if inc.employee else "N/A"
            analyst_name = inc.assigned_analyst.full_name if inc.assigned_analyst else "Unassigned"
            created = inc.created_at.strftime("%Y-%m-%d %H:%M:%S") if inc.created_at else ""
            
            ws.cell(row=row_idx, column=1, value=inc.incident_id).font = data_font
            ws.cell(row=row_idx, column=2, value=inc.title).font = data_font
            ws.cell(row=row_idx, column=3, value=inc.status.value if inc.status else "").font = data_font
            ws.cell(row=row_idx, column=4, value=inc.severity.value if inc.severity else "").font = data_font
            ws.cell(row=row_idx, column=5, value=inc.impact_level).font = data_font
            ws.cell(row=row_idx, column=6, value=emp_name).font = data_font
            ws.cell(row=row_idx, column=7, value=analyst_name).font = data_font
            ws.cell(row=row_idx, column=8, value=created).font = data_font
            
    elif report_type == "alerts":
        headers = ["Alert ID", "Title", "Severity", "Status", "Employee", "Triggered At", "Notes"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
        ws.row_dimensions[row_idx].height = 24
        
        alerts = db.query(Alert).order_by(desc(Alert.triggered_at)).all()
        for alert in alerts:
            row_idx += 1
            emp_name = alert.employee.full_name if alert.employee else "N/A"
            triggered = alert.triggered_at.strftime("%Y-%m-%d %H:%M:%S") if alert.triggered_at else ""
            
            ws.cell(row=row_idx, column=1, value=alert.alert_id).font = data_font
            ws.cell(row=row_idx, column=2, value=alert.title).font = data_font
            ws.cell(row=row_idx, column=3, value=alert.severity.value if alert.severity else "").font = data_font
            ws.cell(row=row_idx, column=4, value=alert.status.value if alert.status else "").font = data_font
            ws.cell(row=row_idx, column=5, value=emp_name).font = data_font
            ws.cell(row=row_idx, column=6, value=triggered).font = data_font
            ws.cell(row=row_idx, column=7, value=alert.notes or "").font = data_font

    elif report_type == "employees":
        headers = ["Employee ID", "Full Name", "Email", "Designation", "Department", "Current Risk Score", "Risk Category"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
        ws.row_dimensions[row_idx].height = 24
        
        from app.api.v1.endpoints.security import get_risk_leaderboard
        leaderboard = get_risk_leaderboard(db, limit=100)
        for entry in leaderboard:
            row_idx += 1
            emp = db.query(Employee).filter(Employee.id == entry.employee_id).first()
            dept_name = emp.department.name if emp and emp.department else "N/A"
            
            ws.cell(row=row_idx, column=1, value=emp.employee_id if emp else "N/A").font = data_font
            ws.cell(row=row_idx, column=2, value=entry.employee_name).font = data_font
            ws.cell(row=row_idx, column=3, value=emp.email if emp else "N/A").font = data_font
            ws.cell(row=row_idx, column=4, value=emp.designation if emp else "N/A").font = data_font
            ws.cell(row=row_idx, column=5, value=dept_name).font = data_font
            ws.cell(row=row_idx, column=6, value=entry.current_score).font = data_font
            ws.cell(row=row_idx, column=7, value=entry.risk_category.value if entry.risk_category else "").font = data_font

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title merge cell
            if cell.row == 1:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"ueba_report_{report_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
