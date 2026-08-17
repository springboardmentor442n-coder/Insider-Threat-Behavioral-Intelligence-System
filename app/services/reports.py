"""Investigation report generation — PDF case files and Excel exports."""

from __future__ import annotations

import io
from datetime import datetime, timezone

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

SEVERITY_COLORS = {
    "CRITICAL": colors.HexColor("#c0392b"),
    "HIGH": colors.HexColor("#d35400"),
    "MEDIUM": colors.HexColor("#b7950b"),
    "LOW": colors.HexColor("#1f6f8b"),
}
INK = colors.HexColor("#1b2733")
MUTED = colors.HexColor("#5a6b7b")
RULE = colors.HexColor("#c8d2dc")


def _styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=base["Title"], fontSize=18,
                                textColor=INK, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=base["Normal"], fontSize=9,
                              textColor=MUTED, spaceAfter=10),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontSize=12,
                             textColor=INK, spaceBefore=12, spaceAfter=6),
        "body": ParagraphStyle("b", parent=base["Normal"], fontSize=9,
                               textColor=INK, alignment=TA_LEFT, leading=13),
        "small": ParagraphStyle("sm", parent=base["Normal"], fontSize=8,
                                textColor=MUTED, leading=11),
    }


def _table(data, col_widths, header_bg="#233648", align_right=()):
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_bg)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.4, RULE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5f8")]),
    ]
    for col in align_right:
        style.append(("ALIGN", (col, 1), (col, -1), "RIGHT"))
    tbl.setStyle(TableStyle(style))
    return tbl


def build_investigation_pdf(case: dict, analyst: str = "unknown") -> bytes:
    """Render a forensic case file for one user-day."""
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title=f"Investigation Report {case['user']}",
        author="Insider Threat Behavioral Intelligence System",
    )

    user = case["user"]
    severity = case.get("severity", "LOW")
    identity = case.get("identity") or {}
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    flow = [
        Paragraph("Insider Threat Investigation Case File", st["title"]),
        Paragraph(
            f"Subject <b>{user}</b> &nbsp;·&nbsp; Observation date {case['day']} "
            f"&nbsp;·&nbsp; Generated {generated} by <b>{analyst}</b>",
            st["sub"],
        ),
    ]

    # --- verdict banner ---------------------------------------------------
    banner = Table(
        [[
            Paragraph(f"<font color='white' size='11'><b>{severity} RISK</b></font>",
                      st["body"]),
            Paragraph(
                f"<font color='white' size='11'>Composite score "
                f"<b>{case['risk_score']:.1f}</b> / 100</font>", st["body"]),
            Paragraph(
                f"<font color='white' size='9'>UEBA {case['ueba_score']:.1f} &nbsp;|&nbsp; "
                f"ML probability {case['ml_probability']:.3f}</font>", st["body"]),
        ]],
        colWidths=[45 * mm, 60 * mm, 69 * mm],
    )
    banner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), SEVERITY_COLORS.get(severity, INK)),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    flow += [banner, Spacer(1, 8)]

    # --- subject ----------------------------------------------------------
    flow.append(Paragraph("1. Subject Identity", st["h2"]))
    flow.append(_table(
        [["Field", "Value"],
         ["User ID", user],
         ["Name", identity.get("name", "Unknown")],
         ["Role", identity.get("role", "Unknown")],
         ["Department", identity.get("department", "Unknown")],
         ["Team", identity.get("team", "Unknown")],
         ["Baseline history", f"{case.get('baseline_days', 0)} observed days"],
         ["Ground-truth label",
          "CONFIRMED INSIDER" if case.get("confirmed_insider") else "Not confirmed"]],
        [45 * mm, 129 * mm]))

    # --- risk breakdown ---------------------------------------------------
    flow.append(Paragraph("2. Weighted Risk Contributions", st["h2"]))
    flow.append(Paragraph(
        "The UEBA engine scores five weighted behavioural indicators against the "
        "subject's own historical baseline. Points shown are that indicator's "
        "contribution to the 0-100 behavioural score.", st["small"]))
    flow.append(Spacer(1, 4))
    rows = [["Indicator", "Weight", "Intensity", "Points"]]
    for c in case.get("risk_contributions", []):
        rows.append([c["label"], f"{c['weight']:.0f}x",
                     f"{c['intensity']:.2f}", f"{c['points']:.1f}"])
    flow.append(_table(rows, [84 * mm, 26 * mm, 32 * mm, 32 * mm], align_right=(1, 2, 3)))

    # --- deviations -------------------------------------------------------
    flow.append(Paragraph("3. Behavioural Deviation Analysis", st["h2"]))
    rows = [["Feature", "Observed", "Baseline", "Deviation (σ)", "Flag"]]
    for d in case.get("deviations", [])[:12]:
        rows.append([
            d["label"], f"{d['observed']:.1f}", f"{d['baseline']:.2f}",
            f"{d['deviation_sigma']:+.2f}", "ANOMALY" if d["anomalous"] else "",
        ])
    flow.append(_table(rows, [64 * mm, 26 * mm, 26 * mm, 30 * mm, 28 * mm],
                       align_right=(1, 2, 3)))

    flow.append(PageBreak())

    # --- explainability ---------------------------------------------------
    explanation = case.get("explanation", {})
    method = "SHAP" if explanation.get("method") == "shap" else "Model importance"
    flow.append(Paragraph(f"4. Model Explainability ({method})", st["h2"]))
    flow.append(Paragraph(
        "Signed feature attributions for this prediction. Positive values pushed "
        "the model toward an insider classification.", st["small"]))
    flow.append(Spacer(1, 4))
    rows = [["Feature", "Value", "Attribution", "Effect"]]
    for f in explanation.get("features", []):
        rows.append([f["label"], f"{f['value']:.1f}",
                     f"{f['contribution']:+.4f}", f["direction"]])
    flow.append(_table(rows, [64 * mm, 26 * mm, 34 * mm, 50 * mm], align_right=(1, 2)))

    # --- peak days --------------------------------------------------------
    flow.append(Paragraph("5. Highest-Risk Days for This Subject", st["h2"]))
    rows = [["Date", "Risk Score", "Severity"]]
    for p in case.get("peak_days", []):
        rows.append([p["day"], f"{p['risk_score']:.1f}", p["severity"]])
    flow.append(_table(rows, [58 * mm, 58 * mm, 58 * mm], align_right=(1,)))

    # --- evidence ---------------------------------------------------------
    flow.append(Paragraph("6. Observed Activity — Evidence Summary", st["h2"]))
    observed = case.get("observed", {})
    labels = {d["feature"]: d["label"] for d in case.get("deviations", [])}
    rows = [["Activity Metric", "Count"]]
    for k, v in observed.items():
        rows.append([labels.get(k, k), f"{v:,.0f}"])
    flow.append(_table(rows, [110 * mm, 64 * mm], align_right=(1,)))

    flow += [
        Spacer(1, 14),
        Paragraph(
            "This report was produced automatically by the Insider Threat Behavioral "
            "Intelligence System. Risk scores are decision-support signals derived from "
            "behavioural modelling, not determinations of misconduct, and must be "
            "corroborated before any action is taken.",
            st["small"],
        ),
    ]

    doc.build(flow)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="233648")
SEVERITY_FILLS = {
    "CRITICAL": PatternFill("solid", fgColor="F5B7B1"),
    "HIGH": PatternFill("solid", fgColor="F8C471"),
    "MEDIUM": PatternFill("solid", fgColor="F9E79F"),
    "LOW": PatternFill("solid", fgColor="D6EAF8"),
}


def _autosize(ws, df, max_width=42):
    for i, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 3, max_width)


def _style_sheet(ws, df):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    _autosize(ws, df)

    if "severity" in df.columns:
        col_idx = list(df.columns).index("severity") + 1
        for row in range(2, len(df) + 2):
            cell = ws.cell(row=row, column=col_idx)
            fill = SEVERITY_FILLS.get(str(cell.value))
            if fill:
                cell.fill = fill


def build_excel(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Write a multi-sheet, styled workbook."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            safe = name[:31]
            if df.empty:
                df = pd.DataFrame({"info": ["no rows matched the current filter"]})
            df.to_excel(writer, sheet_name=safe, index=False)
            _style_sheet(writer.sheets[safe], df)
    return buf.getvalue()
