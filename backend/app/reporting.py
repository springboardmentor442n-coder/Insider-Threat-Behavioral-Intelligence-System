"""Report rendering framework. Spec Module 12.

The problem with "five report types times two formats" is that it invites ten
copy-pasted generators that drift apart. This module avoids that with one idea: a
report is a structured DOCUMENT - a title plus a list of SECTIONS - and each
format has ONE renderer that knows how to draw any document.

    report data function  ->  ReportDocument  ->  render_pdf()  -> bytes
                                              \\-> render_xlsx() -> bytes

So the data for a report is written once (in reports_data.py), and adding a report
never touches the renderers. A section is one of a small, fixed vocabulary:

    SummarySection  - labelled key/value facts (the headline numbers)
    TableSection    - a header row plus data rows
    TextSection     - a narrative paragraph

That vocabulary is deliberately small. A report that needs something outside it is
usually a report trying to be a dashboard, and the two are different things.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime


# ---------------------------------------------------------------------------
# THE DOCUMENT MODEL
# ---------------------------------------------------------------------------

@dataclass
class SummarySection:
    """Labelled facts - the numbers a reader should see first."""
    title: str
    items: list[tuple[str, str]]  # (label, value)


@dataclass
class TableSection:
    """A table: one header row, then data rows. Rows are lists of strings."""
    title: str
    headers: list[str]
    rows: list[list[str]]
    note: str | None = None


@dataclass
class TextSection:
    """A narrative paragraph - context, caveats, methodology."""
    title: str
    body: str


Section = SummarySection | TableSection | TextSection


@dataclass
class ReportDocument:
    """A whole report: identity, a subtitle, and an ordered list of sections."""
    title: str
    subtitle: str
    sections: list[Section] = field(default_factory=list)
    generated_at: datetime = field(default_factory=datetime.utcnow)
    generated_by: str = ""

    def add(self, section: Section) -> "ReportDocument":
        self.sections.append(section)
        return self


# ---------------------------------------------------------------------------
# PALETTE - matches the app's "forensic console" look so a printed report and
# the screen it came from are recognisably the same product.
# ---------------------------------------------------------------------------

_INK = "#0a0e14"
_SLATE = "#334155"
_SIGNAL = "#0891b2"      # a darker cyan than the screen's, for print legibility
_LINE = "#cbd5e1"
_HEADER_BG = "#0f172a"
_ZEBRA = "#f1f5f9"

_SEV_COLORS = {
    "critical": "#dc2626",
    "high": "#ea580c",
    "medium": "#d97706",
    "low": "#0891b2",
    "informational": "#64748b",
}


# ===========================================================================
# PDF RENDERER
# ===========================================================================

def render_pdf(doc: ReportDocument) -> bytes:
    """Render a ReportDocument to PDF bytes with reportlab's platypus layout."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=20 * mm, bottomMargin=18 * mm,
        leftMargin=18 * mm, rightMargin=18 * mm,
        title=doc.title, author="ITBIS",
    )

    styles = getSampleStyleSheet()
    h_title = ParagraphStyle("t", parent=styles["Title"], fontName="Helvetica-Bold",
                             fontSize=20, textColor=colors.HexColor(_INK), spaceAfter=2)
    h_sub = ParagraphStyle("s", parent=styles["Normal"], fontName="Helvetica",
                           fontSize=10, textColor=colors.HexColor(_SLATE), spaceAfter=2)
    h_meta = ParagraphStyle("m", parent=styles["Normal"], fontName="Helvetica",
                            fontSize=8, textColor=colors.HexColor(_SLATE))
    h_sec = ParagraphStyle("sec", parent=styles["Heading2"], fontName="Helvetica-Bold",
                           fontSize=12, textColor=colors.HexColor(_SIGNAL),
                           spaceBefore=14, spaceAfter=6)
    body = ParagraphStyle("b", parent=styles["Normal"], fontName="Helvetica",
                          fontSize=9.5, textColor=colors.HexColor(_INK), leading=14,
                          alignment=TA_LEFT)
    note = ParagraphStyle("n", parent=styles["Normal"], fontName="Helvetica-Oblique",
                          fontSize=8, textColor=colors.HexColor(_SLATE), spaceBefore=3)

    story = []
    story.append(Paragraph(doc.title, h_title))
    story.append(Paragraph(doc.subtitle, h_sub))
    meta = f"Generated {doc.generated_at:%Y-%m-%d %H:%M UTC}"
    if doc.generated_by:
        meta += f" &middot; by {doc.generated_by}"
    story.append(Paragraph(meta, h_meta))
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor(_SIGNAL)))

    for sec in doc.sections:
        if isinstance(sec, SummarySection):
            story.append(Paragraph(sec.title, h_sec))
            data = [[Paragraph(f"<b>{v}</b>", body), Paragraph(lbl, body)]
                    for lbl, v in sec.items]
            t = Table(data, colWidths=[40 * mm, None])
            t.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor(_LINE)),
            ]))
            story.append(t)

        elif isinstance(sec, TableSection):
            story.append(Paragraph(sec.title, h_sec))
            header = [Paragraph(f"<b>{h}</b>", ParagraphStyle(
                "th", parent=body, textColor=colors.white, fontSize=8.5)) for h in sec.headers]
            rows = [[Paragraph(str(c), ParagraphStyle("td", parent=body, fontSize=8.5))
                     for c in r] for r in sec.rows]
            t = Table([header] + rows, repeatRows=1)
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_HEADER_BG)),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(_LINE)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
            for i in range(1, len(rows) + 1):
                if i % 2 == 0:
                    style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(_ZEBRA)))
            t.setStyle(TableStyle(style))
            story.append(t)
            if sec.note:
                story.append(Paragraph(sec.note, note))

        elif isinstance(sec, TextSection):
            story.append(Paragraph(sec.title, h_sec))
            story.append(Paragraph(sec.body, body))

    pdf.build(story)
    return buf.getvalue()


# ===========================================================================
# EXCEL RENDERER
# ===========================================================================

def render_xlsx(doc: ReportDocument) -> bytes:
    """Render a ReportDocument to an .xlsx workbook.

    Layout choice: summaries and text go on an 'Overview' sheet; each table gets
    its OWN sheet, because a table on its own tab is what makes an Excel export
    actually usable for filtering and pivoting - the whole reason to offer Excel
    alongside PDF.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"

    ink = Font(name="Arial", color="0A0E14")
    bold = Font(name="Arial", bold=True, color="0A0E14")
    title_font = Font(name="Arial", bold=True, size=16, color="0A0E14")
    sub_font = Font(name="Arial", size=10, color="334155")
    sec_font = Font(name="Arial", bold=True, size=12, color="0891B2")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F172A")
    zebra_fill = PatternFill("solid", fgColor="F1F5F9")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    r = 1
    overview.cell(r, 1, doc.title).font = title_font
    r += 1
    overview.cell(r, 1, doc.subtitle).font = sub_font
    r += 1
    meta = f"Generated {doc.generated_at:%Y-%m-%d %H:%M UTC}"
    if doc.generated_by:
        meta += f" by {doc.generated_by}"
    overview.cell(r, 1, meta).font = sub_font
    r += 2

    for sec in doc.sections:
        if isinstance(sec, SummarySection):
            overview.cell(r, 1, sec.title).font = sec_font
            r += 1
            for lbl, val in sec.items:
                overview.cell(r, 1, lbl).font = ink
                c = overview.cell(r, 2, _coerce(val))
                c.font = bold
                r += 1
            r += 1

        elif isinstance(sec, TextSection):
            overview.cell(r, 1, sec.title).font = sec_font
            r += 1
            cell = overview.cell(r, 1, sec.body)
            cell.font = ink
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            overview.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            r += 2

        elif isinstance(sec, TableSection):
            # each table on its own sheet
            ws = wb.create_sheet(title=_sheet_name(sec.title, wb))
            for j, h in enumerate(sec.headers, start=1):
                c = ws.cell(1, j, h)
                c.font = header_font
                c.fill = header_fill
                c.border = border
                c.alignment = Alignment(horizontal="left", vertical="center")
            for i, row in enumerate(sec.rows, start=2):
                for j, val in enumerate(row, start=1):
                    c = ws.cell(i, j, _coerce(val))
                    c.font = ink
                    c.border = border
                    if i % 2 == 1:
                        c.fill = zebra_fill
            # autosize-ish: width from the longest cell in each column
            for j in range(1, len(sec.headers) + 1):
                longest = len(str(sec.headers[j - 1]))
                for row in sec.rows:
                    if j - 1 < len(row):
                        longest = max(longest, len(str(row[j - 1])))
                ws.column_dimensions[get_column_letter(j)].width = min(longest + 3, 50)
            ws.freeze_panes = "A2"
            if sec.note:
                nr = len(sec.rows) + 3
                ws.cell(nr, 1, sec.note).font = sub_font

            # reference the table sheet from the overview
            overview.cell(r, 1, f"See sheet: {ws.title}").font = ink
            r += 1

    # overview column widths
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _coerce(val: str):
    """Store numbers as numbers so Excel can sum/sort them, not as text."""
    if isinstance(val, (int, float)):
        return val
    s = str(val)
    # a percentage or sigma string stays text; a bare number becomes numeric
    try:
        if s.replace(".", "", 1).replace("-", "", 1).isdigit():
            return float(s) if "." in s else int(s)
    except (ValueError, AttributeError):
        pass
    return s


def _sheet_name(title: str, wb) -> str:
    """Excel sheet names: <=31 chars, no []:*?/\\ and must be unique."""
    clean = "".join(c for c in title if c not in "[]:*?/\\")[:28]
    name = clean or "Sheet"
    i = 2
    existing = set(wb.sheetnames)
    base = name
    while name in existing:
        name = f"{base[:26]}_{i}"
        i += 1
    return name
